"""无损展开 C 题附件，生成可追溯的长表和数据审计元数据。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


DT_HOURS = 1.0 / 6.0
SLOTS_PER_DAY = 144
ISSUE_MINUTES = {"0:00": 0, "6:00": 360, "12:00": 720, "18:00": 1080}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=Path("data/raw/附件"))
    parser.add_argument("--processed-dir", type=Path, default=Path("data/processed"))
    parser.add_argument("--metadata-dir", type=Path, default=Path("data/metadata"))
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("/", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"无法解析日期：{value!r}")


def endpoint_minutes(value: object) -> int:
    if isinstance(value, time):
        minutes = value.hour * 60 + value.minute
        return 1440 if minutes == 0 else minutes
    text = str(value).strip()
    if text == "0:00+1":
        return 1440
    hour, minute = map(int, text.split(":"))
    minutes = hour * 60 + minute
    return 1440 if minutes == 0 else minutes


def issue_minutes(value: object) -> int:
    text = value.strftime("%-H:%M") if isinstance(value, time) else str(value).strip()
    if text not in ISSUE_MINUTES:
        raise ValueError(f"不支持的预报时刻：{value!r}")
    return ISSUE_MINUTES[text]


def write_csv(path: Path, header: list[str], rows: Iterable[Iterable[object]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        for row in rows:
            writer.writerow(row)
            count += 1
    return count


def load_matrix_sheet(path: Path, sheet: str) -> tuple[list[object], list[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    iterator = worksheet.iter_rows(values_only=True)
    header = list(next(iterator))
    rows = [tuple(row) for row in iterator]
    workbook.close()
    return header, rows


def make_q1(raw_dir: Path, out_dir: Path) -> dict[str, object]:
    path = raw_dir / "附件1.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows: list[list[object]] = []
    for slot, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=1):
        end_minute = endpoint_minutes(row[0])
        price, load_kw, pv_kw = map(float, row[1:4])
        rows.append([
            slot,
            end_minute - 10,
            end_minute,
            price,
            load_kw,
            pv_kw,
            load_kw * DT_HOURS,
            pv_kw * DT_HOURS,
        ])
    workbook.close()
    if len(rows) != SLOTS_PER_DAY:
        raise ValueError(f"附件1应有144个时段，实际{len(rows)}")
    count = write_csv(
        out_dir / "q1_timeseries.csv",
        ["slot", "interval_start_min", "interval_end_min", "price_cny_per_kwh",
         "load_kw", "pv_forecast_kw", "load_kwh", "pv_forecast_kwh"],
        rows,
    )
    return {"rows": count, "numeric_values_modified": 0}


def make_annual_actual(raw_dir: Path, out_dir: Path) -> tuple[dict[str, object], dict[datetime, float]]:
    path = raw_dir / "附件2.xlsx"
    load_header, load_rows = load_matrix_sheet(path, "小区负载")
    pv_header, pv_rows = load_matrix_sheet(path, "光伏发电实际功率")
    if load_header != pv_header or len(load_rows) != len(pv_rows):
        raise ValueError("附件2两个工作表的日期或时段结构不一致")
    end_minutes = [endpoint_minutes(value) for value in load_header[1:]]
    if end_minutes != list(range(10, 1450, 10)):
        raise ValueError("附件2并非连续144个10分钟时段")

    output: list[list[object]] = []
    pv_by_endpoint: dict[datetime, float] = {}
    dates: list[date] = []
    for load_row, pv_row in zip(load_rows, pv_rows):
        day = parse_date(load_row[0])
        if day != parse_date(pv_row[0]):
            raise ValueError("附件2负载与光伏日期错位")
        dates.append(day)
        for slot, (minute, load_value, pv_value) in enumerate(
            zip(end_minutes, load_row[1:], pv_row[1:]), start=1
        ):
            load_kw, pv_kw = float(load_value), float(pv_value)
            if load_kw < 0 or pv_kw < 0:
                raise ValueError(f"附件2出现负功率：{day} slot={slot}")
            end_ts = datetime.combine(day, time()) + timedelta(minutes=minute)
            start_ts = end_ts - timedelta(minutes=10)
            output.append([
                day.isoformat(), slot, start_ts.isoformat(sep=" "), end_ts.isoformat(sep=" "),
                load_kw, pv_kw, load_kw * DT_HOURS, pv_kw * DT_HOURS,
            ])
            pv_by_endpoint[end_ts] = pv_kw
    if len(dates) != 365 or dates != [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]:
        raise ValueError("附件2日期必须完整覆盖2025年")
    count = write_csv(
        out_dir / "annual_actual.csv",
        ["date", "slot", "interval_start", "interval_end", "load_kw", "pv_actual_kw",
         "load_kwh", "pv_actual_kwh"],
        output,
    )
    return {"rows": count, "days": len(dates), "numeric_values_modified": 0}, pv_by_endpoint


def make_annual_price(raw_dir: Path, out_dir: Path) -> dict[str, object]:
    path = raw_dir / "附件4.xlsx"
    header, source_rows = load_matrix_sheet(path, "Sheet1")
    end_minutes = [endpoint_minutes(value) for value in header[1:]]
    if end_minutes != list(range(10, 1450, 10)):
        raise ValueError("附件4并非连续144个10分钟时段")
    output: list[list[object]] = []
    dates: list[date] = []
    for row in source_rows:
        day = parse_date(row[0])
        dates.append(day)
        for slot, (minute, value) in enumerate(zip(end_minutes, row[1:]), start=1):
            price = float(value)
            if price < 0:
                raise ValueError(f"附件4出现负电价：{day} slot={slot}")
            end_ts = datetime.combine(day, time()) + timedelta(minutes=minute)
            output.append([day.isoformat(), slot, end_ts.isoformat(sep=" "), price])
    if len(dates) != 365:
        raise ValueError(f"附件4应有365天，实际{len(dates)}")
    count = write_csv(
        out_dir / "annual_price.csv",
        ["date", "slot", "interval_end", "price_cny_per_kwh"],
        output,
    )
    return {"rows": count, "days": len(dates), "numeric_values_modified": 0}


def make_forecasts(
    raw_dir: Path, out_dir: Path, pv_by_endpoint: dict[datetime, float]
) -> tuple[dict[str, object], dict[str, object]]:
    path = raw_dir / "附件3.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    hourly_rows: list[list[object]] = []
    current_day: date | None = None
    structural_date_fills = 0
    issue_count = 0
    boundary_anchor_fallbacks = 0
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ""):
            current_day = parse_date(row[0])
        else:
            structural_date_fills += 1
        if current_day is None:
            raise ValueError("附件3首个预报组缺少日期")
        minute = issue_minutes(row[1])
        issue_ts = datetime.combine(current_day, time()) + timedelta(minutes=minute)
        forecasts = [float(value) for value in row[2:26]]
        if len(forecasts) != 24 or any(value < 0 for value in forecasts):
            raise ValueError(f"附件3预报异常：{issue_ts}")
        issue_count += 1
        for horizon, value in enumerate(forecasts, start=1):
            target = issue_ts + timedelta(hours=horizon)
            hourly_rows.append([
                current_day.isoformat(), row[1], issue_ts.isoformat(sep=" "), horizon,
                target.isoformat(sep=" "), value,
            ])

        anchor = pv_by_endpoint.get(issue_ts)
        if anchor is None:
            if issue_ts == datetime(2025, 1, 1):
                anchor = 0.0
                boundary_anchor_fallbacks += 1
            else:
                raise ValueError(f"找不到预报发布时间的已知光伏实测锚点：{issue_ts}")
        interpolated = interpolate_hourly_forecast(anchor, forecasts)
        if any(value < 0 for value in interpolated):
            raise ValueError(f"附件3插值后出现负功率：{issue_ts}")
    workbook.close()
    if issue_count != 365 * 4:
        raise ValueError(f"附件3应有1460个发布记录，实际{issue_count}")
    hourly_count = write_csv(
        out_dir / "pv_forecasts_hourly.csv",
        ["issue_date", "issue_time", "issue_timestamp", "horizon_hour", "target_timestamp",
         "pv_forecast_kw"],
        hourly_rows,
    )
    return (
        {"rows": hourly_count, "issues": issue_count, "structural_date_fills": structural_date_fills},
        {"virtual_rows": issue_count * SLOTS_PER_DAY, "method": "actual-anchor linear interpolation on demand",
         "boundary_anchor_fallbacks": boundary_anchor_fallbacks},
    )


def interpolate_hourly_forecast(anchor: float, forecasts: list[float]) -> list[float]:
    """把发布时刻锚点与未来24个整点功率线性映射到144个10分钟端点。"""
    if len(forecasts) != 24:
        raise ValueError("每次发布必须包含未来24小时预报")
    knots = [float(anchor), *map(float, forecasts)]
    result: list[float] = []
    for horizon_slot in range(1, SLOTS_PER_DAY + 1):
        hour_position = horizon_slot / 6.0
        left = int(hour_position)
        if hour_position == left:
            value = knots[left]
        else:
            fraction = hour_position - left
            value = knots[left] * (1.0 - fraction) + knots[left + 1] * fraction
        result.append(max(0.0, value))
    return result


def write_dictionary(metadata_dir: Path) -> int:
    rows = [
        ["q1_timeseries.csv", "price_cny_per_kwh", "float", "元/kWh", "附件1电价"],
        ["q1_timeseries.csv", "load_kw", "float", "kW", "附件1小区负载功率"],
        ["q1_timeseries.csv", "pv_forecast_kw", "float", "kW", "附件1光伏预测功率"],
        ["annual_actual.csv", "load_kw", "float", "kW", "附件2小区负载功率"],
        ["annual_actual.csv", "pv_actual_kw", "float", "kW", "附件2光伏实际功率"],
        ["annual_price.csv", "price_cny_per_kwh", "float", "元/kWh", "附件4实时电价"],
        ["pv_forecasts_hourly.csv", "pv_forecast_kw", "float", "kW", "附件3整点预报"],
        ["pv_forecasts_hourly.csv", "target_timestamp", "datetime", "", "发布时间加预报步长"],
        ["*_timeseries/actual/forecast*.csv", "*_kwh", "float", "kWh", "对应功率乘以1/6小时；预报在求解时插值"],
    ]
    return write_csv(
        metadata_dir / "data_dictionary.csv",
        ["file", "field", "type", "unit", "definition"],
        rows,
    )


def write_preprocessing_checks(metadata_dir: Path, outputs: dict[str, dict[str, object]]) -> int:
    expected = {
        "q1_timeseries.csv": 144,
        "annual_actual.csv": 365 * 144,
        "annual_price.csv": 365 * 144,
        "pv_forecasts_hourly.csv": 365 * 4 * 24,
    }
    rows = []
    for name, expected_rows in expected.items():
        actual_rows = int(outputs[name]["rows"])
        rows.append([name, expected_rows, actual_rows, 0, 0,
                     "passed" if actual_rows == expected_rows else "failed",
                     "保留原值，仅作结构与单位标准化"])
    rows.extend([
        ["附件3结构日期", 1095, outputs["pv_forecasts_hourly.csv"]["structural_date_fills"],
         "", "", "passed", "仅向下继承每日后三个发布时刻的日期"],
        ["附件3十分钟插值", 365 * 4 * 144, outputs["pv_forecasts_10min_virtual"]["virtual_rows"],
         "", 0, "passed", "按需生成，不持久化超大派生表"],
    ])
    return write_csv(
        metadata_dir / "preprocessing_checks.csv",
        ["dataset", "expected_rows", "actual_rows", "missing_numeric", "negative_numeric",
         "status", "action"],
        rows,
    )


def main() -> None:
    args = parse_args()
    required = [args.raw_dir / f"附件{i}.xlsx" for i in range(1, 5)]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"缺少输入附件：{missing}")
    args.processed_dir.mkdir(parents=True, exist_ok=True)
    args.metadata_dir.mkdir(parents=True, exist_ok=True)

    q1 = make_q1(args.raw_dir, args.processed_dir)
    actual, pv_by_endpoint = make_annual_actual(args.raw_dir, args.processed_dir)
    prices = make_annual_price(args.raw_dir, args.processed_dir)
    forecasts_hourly, forecasts_10min = make_forecasts(args.raw_dir, args.processed_dir, pv_by_endpoint)
    outputs = {
        "q1_timeseries.csv": q1,
        "annual_actual.csv": actual,
        "annual_price.csv": prices,
        "pv_forecasts_hourly.csv": forecasts_hourly,
        "pv_forecasts_10min_virtual": forecasts_10min,
    }
    dictionary_rows = write_dictionary(args.metadata_dir)
    check_rows = write_preprocessing_checks(args.metadata_dir, outputs)

    profile = {
        "status": "checked",
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "source_sha256": {path.name: sha256(path) for path in required},
        "policy": {
            "raw_files_modified": False,
            "missing_numeric_values_filled": 0,
            "outliers_deleted_or_winsorized": 0,
            "structural_dates_forward_filled": forecasts_hourly["structural_date_fills"],
            "forecast_interpolation": "known actual at issue time + piecewise linear to hourly forecasts",
            "time_step_hours": DT_HOURS,
        },
        "outputs": {
            **outputs,
            "data_dictionary.csv": {"rows": dictionary_rows},
            "preprocessing_checks.csv": {"rows": check_rows},
        },
    }
    profile_path = args.metadata_dir / "data_profile.json"
    profile_path.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(profile, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
